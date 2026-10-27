r"""Final results & comparison report for the faculty: every model of every modality, evaluated the same way.

Reads project/reports/*_metrics.json (end-to-end test results of each trained hybrid model),
project/reports/protocol/*.json (faculty protocol: 80:20 / 70:30 +/- SMOTE, 10-fold CV +/- SMOTE) and
project/reports/face_tta_ensemble_metrics.json, draws comparison charts and confusion matrices, and writes
project/reports/final/Final_Results_Multimodal_Emotion.docx

Usage (repo root): project\.venv-train\Scripts\python project\code\make_final_report.py
"""
import json
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
from docx.shared import Pt  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402

sys.path.insert(0, str(Path(__file__).parent))
import make_review2 as mr  # noqa: E402  (shared Doc class, team, title)

ROOT = Path(__file__).resolve().parents[2]
REP = ROOT / "project" / "reports"
PROT = REP / "protocol"
OUT = REP / "final"
FIG = OUT / "fig"

NAMES = {
    "face": {"cnn": "CNN (baseline)", "resnet18": "ResNet-18 (baseline)", "effnet_cbam": "EfficientNetB3 + CBAM",
             "resnet50_vit": "ResNet50 + ViT", "densenet_se": "DenseNet121 + SE",
             "mobilenet_attn": "MobileNetV3 + Attention", "cnn_bilstm": "CNN + BiLSTM",
             "stack": "Stacked features (EffNet + DenseNet)"},
    "audio": {"wav2vec2_bilstm_attn": "wav2vec2 + BiLSTM + Attn", "cnn_bilstm_attn": "CNN + BiLSTM + Attn",
              "hubert_transformer": "HuBERT + Transformer", "ast_mlp": "AST + MLP", "cnn_gru": "CNN + GRU"},
    "text": {"roberta_bilstm_attn": "RoBERTa + BiLSTM + Attn", "bert_bilstm_attn": "BERT + BiLSTM + Attn",
             "deberta_attn": "DeBERTa + Attn", "distilbert_cnn": "DistilBERT + CNN",
             "bert_cnn_bilstm": "BERT + CNN + BiLSTM"},
    "wesad": {"cnn_lstm": "CNN-LSTM", "cnn_bilstm": "CNN-BiLSTM", "cnn_gru": "1D-CNN-GRU",
              "ae_mlp": "Autoencoder-MLP", "transformer_bilstm": "Transformer-BiLSTM"},
}
SETTINGS = [("holdout_80_20_no_smote", "80:20"), ("holdout_80_20_smote", "80:20 + SMOTE"),
            ("holdout_70_30_no_smote", "70:30"), ("holdout_70_30_smote", "70:30 + SMOTE"),
            ("cv10_no_smote", "10-fold"), ("cv10_smote", "10-fold + SMOTE")]
BLUE, ORANGE, GREEN, GREY = "#2e6da4", "#e08a1e", "#3c8d3c", "#9aa5b1"


def load(p):
    return json.loads(Path(p).read_text())


def e2e(mod, model, dataset=None):
    f = REP / (f"{mod}_{model}_{dataset}_metrics.json" if dataset else f"{mod}_{model}_metrics.json")
    return load(f) if f.exists() else None


def prot(stem):
    f = PROT / f"{stem}.json"
    return load(f) if f.exists() else None


def macro(m, key):
    return np.mean([v[key] for v in m["per_class"].values()])


def timing(m):
    t = m["inference_time"]
    cpu = next(v for k, v in t.items() if k.startswith("cpu"))
    gpu = next((v for k, v in t.items() if k.startswith("cuda")), None)
    return cpu, gpu


def fmt_setting(r, key):
    v = r.get(key)
    if v is None:
        return "–"
    if key.startswith("cv10"):
        return f"{v['accuracy_mean']:.3f} / {v['macro_f1_mean']:.3f} ± {v['macro_f1_std']:.3f}"
    return f"{v['accuracy']:.3f} / {v['macro_f1']:.3f}"


def save(fig, name):
    FIG.mkdir(parents=True, exist_ok=True)
    p = FIG / f"{name}.png"
    fig.savefig(p, bbox_inches="tight", facecolor="white", dpi=160)
    plt.close(fig)
    return p


def compare_chart(name, title, labels, test_f1, cv_f1, cv_std):
    fig, ax = plt.subplots(figsize=(10, 3.8))
    x = np.arange(len(labels))
    b1 = ax.bar(x - 0.2, test_f1, 0.4, label="Held-out test set (end-to-end model)", color=BLUE)
    b2 = ax.bar(x + 0.2, cv_f1, 0.4, yerr=cv_std, capsize=3, label="10-fold CV + SMOTE (protocol)", color=ORANGE)
    for b in (b1, b2):
        ax.bar_label(b, fmt="%.3f", fontsize=7.5, padding=2)
    ax.set_xticks(x, labels, rotation=20, ha="right", fontsize=8.5)
    ax.set_ylabel("Macro-F1")
    lo = np.nanmin(np.r_[test_f1, cv_f1]) - 0.1
    ax.set_ylim(max(0, lo), 1.0)
    ax.legend(fontsize=8, loc="upper left", ncol=2)
    ax.set_title(title, fontweight="bold")
    ax.spines[["top", "right"]].set_visible(False)
    return save(fig, name)


def time_chart(name, title, labels, cpu, gpu, unit):
    fig, ax = plt.subplots(figsize=(10, 3.2))
    x = np.arange(len(labels))
    b1 = ax.bar(x - 0.2, cpu, 0.4, label="CPU (Ryzen 7 6800H)", color=GREY)
    b2 = ax.bar(x + 0.2, [g or 0 for g in gpu], 0.4, label="GPU (RTX 3050 Laptop)", color=GREEN)
    for b in (b1, b2):
        ax.bar_label(b, fmt="%.1f", fontsize=7.5, padding=2)
    ax.set_xticks(x, labels, rotation=20, ha="right", fontsize=8.5)
    ax.set_ylabel(f"ms per {unit}")
    ax.legend(fontsize=8)
    ax.set_title(title, fontweight="bold")
    ax.spines[["top", "right"]].set_visible(False)
    return save(fig, name)


def cm_figure(name, title, m):
    cm = np.array(m["confusion_matrix"]["rows_true_cols_pred"])
    labels = m["confusion_matrix"]["labels"]
    norm = cm / cm.sum(1, keepdims=True).clip(min=1)
    fig, ax = plt.subplots(figsize=(1.0 * len(labels) + 1.6, 0.85 * len(labels) + 1.2))
    ax.imshow(norm, cmap="Blues", vmin=0, vmax=1)
    for i in range(len(labels)):
        for j in range(len(labels)):
            ax.text(j, i, f"{cm[i, j]}\n{norm[i, j]:.0%}", ha="center", va="center", fontsize=7.5,
                    color="white" if norm[i, j] > 0.55 else "#222")
    ax.set_xticks(range(len(labels)), labels, rotation=35, ha="right")
    ax.set_yticks(range(len(labels)), labels)
    ax.set_xlabel("Predicted")
    ax.set_ylabel("True")
    ax.set_title(title, fontweight="bold", fontsize=10)
    return save(fig, name)


def smote_chart(name, title, rows):
    """rows: (label, f1_no_smote, f1_smote) for the 10-fold setting."""
    fig, ax = plt.subplots(figsize=(10, 3.2))
    x = np.arange(len(rows))
    b1 = ax.bar(x - 0.2, [r[1] for r in rows], 0.4, label="10-fold, no SMOTE", color=GREY)
    b2 = ax.bar(x + 0.2, [r[2] for r in rows], 0.4, label="10-fold, with SMOTE", color=ORANGE)
    for b in (b1, b2):
        ax.bar_label(b, fmt="%.3f", fontsize=7, padding=2)
    ax.set_xticks(x, [r[0] for r in rows], rotation=25, ha="right", fontsize=7.5)
    ax.set_ylabel("Mean macro-F1")
    ax.set_ylim(max(0, min(min(r[1], r[2]) for r in rows) - 0.1), 1.0)
    ax.legend(fontsize=8)
    ax.set_title(title, fontweight="bold")
    ax.spines[["top", "right"]].set_visible(False)
    return save(fig, name)


def modality_section(doc, mod, title, entries, unit, intro, extra=None):
    """entries: list of (display, e2e-metrics, protocol-json, stem-for-figs)."""
    short = title.split(" ", 1)[1]
    doc.h(title, 2)
    doc.p(intro)
    rows = []
    for disp, m, _, _ in entries:
        if m is None:
            rows.append((disp, "–", "–", "–", "–", "–", "–"))
            continue
        cpu, gpu = timing(m)
        rows.append((disp, f"{m['accuracy']:.4f}", f"{macro(m, 'precision'):.4f}", f"{macro(m, 'recall'):.4f}",
                     f"{m['macro_f1']:.4f}", f"{cpu:.1f}", f"{gpu:.1f}" if gpu else "–"))
    best_i = int(np.argmax([m["macro_f1"] if m else -1 for _, m, _, _ in entries]))
    rows[best_i] = tuple([rows[best_i][0] + "  ★"] + list(rows[best_i][1:]))
    doc.table(["Model", "Accuracy", "Precision (macro)", "Recall (macro)", "F1 (macro)", f"CPU ms/{unit}",
               f"GPU ms/{unit}"], rows, f"{short} – held-out test results of the trained models "
              "(★ = best macro-F1)", font=8.5)
    prow = []
    for disp, _, p, _ in entries:
        if p is None:
            continue
        fits = {p[k]["fit"] for k, _ in SETTINGS[:4] if k in p}
        prow.append([disp] + [fmt_setting(p, k) for k, _ in SETTINGS] + [", ".join(sorted(fits))])
    if prow:
        doc.table(["Model"] + [s for _, s in SETTINGS] + ["Fit check"], prow,
                  f"{short} – faculty protocol results (accuracy / macro-F1; 10-fold: mean ± std)", font=7)
    sel = [(d, m, p) for d, m, p, _ in entries if m and p and "cv10_smote" in p]
    if sel:
        doc.img(compare_chart(f"{mod}_compare", f"{short} – macro-F1 comparison", [s[0] for s in sel],
                              [s[1]["macro_f1"] for s in sel], [s[2]["cv10_smote"]["macro_f1_mean"] for s in sel],
                              [s[2]["cv10_smote"]["macro_f1_std"] for s in sel]),
                "Macro-F1 of every model: held-out test vs 10-fold cross-validation with SMOTE")
        sm = [(d, p["cv10_no_smote"]["macro_f1_mean"], p["cv10_smote"]["macro_f1_mean"])
              for d, _, p in sel if "cv10_no_smote" in p]
        if sm:
            doc.img(smote_chart(f"{mod}_smote", f"{short} – effect of SMOTE (10-fold CV)", sm),
                    "10-fold macro-F1 with and without SMOTE")
    tl = [(d, *timing(m)) for d, m, _, _ in entries if m]
    doc.img(time_chart(f"{mod}_time", f"{short} – inference time", [t[0] for t in tl],
                       [t[1] for t in tl], [t[2] for t in tl], unit), f"Inference time per {unit} (batch size 1)")
    bd, bm, _, _ = entries[best_i]
    doc.img(cm_figure(f"{mod}_cm_best", f"Confusion matrix – {bd} (test set)", bm),
            f"Confusion matrix of the best model ({bd}) on the held-out test set", width=4.8)
    if extra:
        extra()
    return entries[best_i]


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    doc = mr.Doc()
    d = doc.d
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("RESULTS AND COMPARATIVE ANALYSIS")
    r.bold, r.font.size = True, Pt(16)
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(mr.TITLE)
    r.bold, r.font.size = True, Pt(14)
    doc.table(["S. No.", "Name", "Registration number"], [(i + 1, n, g) for i, (n, g) in enumerate(mr.TEAM)],
              "Team members", widths=[0.8, 2.6, 2.2], font=11)

    doc.h("1. Experimental Setup", 1)
    doc.bullets([
        ("Hardware: ", "AMD Ryzen 7 6800H CPU, 14 GB RAM, NVIDIA GeForce RTX 3050 Laptop GPU (4 GB); "
         "PyTorch 2.6 with CUDA 12.4, mixed precision."),
        ("Models: ", "20 hybrid deep-learning models (5 per modality, as specified) + 2 face baselines; "
         "text models are trained on two datasets, giving 27 trained networks."),
        ("Evaluation A – held-out test: ", "each model is trained on the training split and tested once on a "
         "test split it never saw (FER-2013 official test set; 15% stratified audio test set; official text test "
         "sets; 3 unseen WESAD subjects)."),
        ("Evaluation B – faculty protocol: ", "80:20 and 70:30 hold-out and stratified 10-fold cross-validation, "
         "each with and without SMOTE on the training data only; early stopping and L1 + L2 regularisation; "
         "over/underfitting check from the train-validation gap. For face, speech and text the protocol runs on "
         "the CSV of deep features of held-out samples; for WESAD the hybrid models are re-trained in every split, "
         "and the splits are by subject (no person appears in both training and test)."),
        ("Metrics: ", "accuracy, macro precision, macro recall, macro-F1 (main metric, robust to class "
         "imbalance), confusion matrix, inference time on CPU and GPU."),
    ])

    doc.h("2. Results per Modality", 1)
    best = {}
    ens = load(REP / "face_tta_ensemble_metrics.json")

    face_models = ["cnn", "resnet18", "effnet_cbam", "resnet50_vit", "densenet_se", "mobilenet_attn", "cnn_bilstm"]
    fe = [(NAMES["face"][m], e2e("face", m), prot(f"face_{m}"), m) for m in face_models]

    def face_extra():
        rows = [(NAMES["face"][m], f"{v['test_plain']['accuracy']:.4f}", f"{v['test_plain']['macro_f1']:.4f}",
                 f"{v['test_tta']['accuracy']:.4f}", f"{v['test_tta']['macro_f1']:.4f}")
                for m, v in ens["tta"].items()]
        e = ens["ensemble"]
        rows.append(("Ensemble: " + " + ".join(NAMES["face"][m] for m in e["members"]) + "  ★", "–", "–",
                     f"{e['test']['accuracy']:.4f}", f"{e['test']['macro_f1']:.4f}"))
        doc.table(["Model", "Accuracy", "Macro-F1", "Accuracy (TTA)", "Macro-F1 (TTA)"], rows,
                  "Face – test-time augmentation (horizontal flip) and validation-selected ensemble", font=8.5)
        doc.p("Averaging the prediction of each image with its mirror image (test-time augmentation) improves "
              "almost every model by up to 1.3 macro-F1 points at no training cost. Averaging the two best models, "
              "chosen on the validation split (never on test), gives the best face result overall: "
              f"{e['test']['accuracy']:.1%} accuracy and {e['test']['macro_f1']:.3f} macro-F1.")
        doc.img(cm_figure("face_cm_ensemble", "Confusion matrix – face ensemble (test set)",
                          {"confusion_matrix": {"labels": list(e["test"]["per_class"]),
                                                "rows_true_cols_pred": e["test"]["confusion_matrix"]}}),
                "Confusion matrix of the face ensemble on the FER-2013 test set", width=4.8)
        st = prot("face_stack")
        if st:
            doc.p(f"Stacked deep features of the two ensemble members also give the best face protocol score: "
                  f"{st['cv10_smote']['macro_f1_mean']:.3f} ± {st['cv10_smote']['macro_f1_std']:.3f} macro-F1 "
                  f"(10-fold CV with SMOTE).")

    best["Face (FER-2013)"] = modality_section(
        doc, "face", "2.1 Facial Expression (FER-2013)", fe, "image",
        "All face models were trained on the 28,709 FER-2013 training images (90% train / 10% validation) and "
        "tested on the 7,178 official test images. Human accuracy on FER-2013 is about 65 ± 5%, and published "
        "single models typically reach 70–76%.", face_extra)

    am = ["wav2vec2_bilstm_attn", "cnn_bilstm_attn", "hubert_transformer", "ast_mlp", "cnn_gru"]
    best["Speech (RAVDESS+TESS+SAVEE)"] = modality_section(
        doc, "audio", "2.2 Speech Emotion (RAVDESS + TESS + SAVEE)",
        [(NAMES["audio"][m], e2e("audio", m), prot(f"audio_{m}"), m) for m in am], "clip",
        "7,520 clips were split 70 / 15 / 15 (stratified by corpus and emotion). The self-supervised encoders "
        "(wav2vec 2.0, HuBERT, AST) are frozen and their features cached; the CNN models learn from log-mel "
        "spectrograms. Note: the split is by clip, so the same speakers (TESS has only two) occur in training and "
        "test – these scores are therefore optimistic for completely new speakers (see Section 4).")

    tm = ["roberta_bilstm_attn", "bert_bilstm_attn", "deberta_attn", "distilbert_cnn", "bert_cnn_bilstm"]
    best["Text (GoEmotions)"] = modality_section(
        doc, "text_go", "2.3 Text Emotion (GoEmotions, 7 classes)",
        [(NAMES["text"][m], e2e("text", m, "goemotions"), prot(f"text_{m}_goemotions"), m) for m in tm], "text",
        "Transformer encoders were fine-tuned end-to-end for 4 epochs (DeBERTa: 2 epochs – the run was "
        "interrupted, and the best checkpoint of epoch 2 is reported). GoEmotions is a hard, noisy Reddit dataset "
        "with very rare 'disgust' and 'fear' classes, so macro-F1 is far below accuracy.")
    best["Text (dair-ai Emotion)"] = modality_section(
        doc, "text_em", "2.4 Text Emotion (dair-ai Emotion, 6 classes)",
        [(NAMES["text"][m], e2e("text", m, "emotion"), prot(f"text_{m}_emotion"), m) for m in tm], "text",
        "Same five architectures fine-tuned on the 16,000 training tweets of dair-ai Emotion and tested on its "
        "2,000-tweet test set.")

    wm = ["cnn_lstm", "cnn_bilstm", "cnn_gru", "ae_mlp", "transformer_bilstm"]
    best["Physiological (WESAD)"] = modality_section(
        doc, "wesad", "2.5 Physiological Signals (WESAD)",
        [(NAMES["wesad"][m], e2e("wesad", m), prot(f"wesad_{m}"), m) for m in wm], "window",
        "Three-class affect recognition (baseline / stress / amusement) from 60-second chest-sensor windows. "
        "Test subjects S2, S9, S16 never appear in training. In the protocol, every split and every fold is "
        "made by subject and the model is re-trained each time, so the scores measure performance on new people.")

    doc.h("3. Best Model per Modality", 1)
    rows = []
    for modality, (disp, m, p, _) in best.items():
        cv = (f"{p['cv10_smote']['macro_f1_mean']:.3f} ± {p['cv10_smote']['macro_f1_std']:.3f}"
              if p and "cv10_smote" in p else "–")
        cpu, gpu = timing(m)
        rows.append((modality, disp, f"{m['accuracy']:.4f}", f"{m['macro_f1']:.4f}", cv,
                     f"{gpu:.1f}" if gpu else f"{cpu:.1f} (CPU)"))
    e = ens["ensemble"]
    rows.insert(1, ("Face – ensemble", " + ".join(NAMES["face"][m] for m in e["members"]),
                    f"{e['test']['accuracy']:.4f}", f"{e['test']['macro_f1']:.4f}", "–", "–"))
    doc.table(["Modality", "Selected model", "Test accuracy", "Test macro-F1", "10-fold macro-F1 (SMOTE)",
               "GPU ms / sample"], rows, "Selected best model for each modality", font=8.5)
    fig, ax = plt.subplots(figsize=(9, 3.4))
    labs = [r[0] for r in rows]
    x = np.arange(len(labs))
    b1 = ax.bar(x - 0.2, [float(r[2]) for r in rows], 0.4, label="Accuracy", color=BLUE)
    b2 = ax.bar(x + 0.2, [float(r[3]) for r in rows], 0.4, label="Macro-F1", color=ORANGE)
    for b in (b1, b2):
        ax.bar_label(b, fmt="%.3f", fontsize=7.5, padding=2)
    ax.set_xticks(x, labs, rotation=20, ha="right", fontsize=8.5)
    ax.set_ylim(0.4, 1.0)
    ax.legend(fontsize=8)
    ax.set_title("Best model per modality – held-out test", fontweight="bold")
    ax.spines[["top", "right"]].set_visible(False)
    doc.img(save(fig, "best_overall"), "Test accuracy and macro-F1 of the selected model for each modality")

    doc.h("4. Discussion", 1)
    doc.h("4.1 Observations", 2)
    doc.bullets([
        ("Attention and channel re-weighting help on faces: ", "DenseNet121+SE and EfficientNetB3+CBAM are the "
         "strongest single face models; the light MobileNetV3+Attention is the fastest on CPU and is the best "
         "choice for real-time use on a laptop without GPU."),
        ("Self-supervised speech encoders: ", "HuBERT+Transformer gives the best speech result even with a frozen "
         "encoder; spectrogram CNN+BiLSTM/GRU models are close behind and much cheaper."),
        ("Text: ", "all five transformer hybrids are within a few macro-F1 points of each other on both datasets; "
         "the larger encoders (RoBERTa, DeBERTa) are slightly better on GoEmotions but 3–5× slower than "
         "DistilBERT."),
        ("Physiological: ", "on the fixed test subjects 1D-CNN-GRU is best (0.769 macro-F1). Across 10 subject-wise "
         "folds all five models lie between 0.66 and 0.71 macro-F1 (Transformer-BiLSTM highest), a smaller "
         "difference than the person-to-person spread (± 0.15; single subjects range from about 0.5 to 1.0), so "
         "on WESAD the choice of subject matters more than the choice of architecture."),
        ("Effect of SMOTE (10-fold macro-F1, with minus without): ", "speech ±0.002 (no effect – the data are "
         "almost balanced); face −0.007 to −0.019; dair-ai Emotion +0.004 to +0.025; GoEmotions −0.030 to −0.046; "
         "WESAD −0.058 to +0.035 (within the subject-to-subject spread). SMOTE helps a little where the minority "
         "classes are moderately small (dair-ai Emotion) but hurts on GoEmotions, where synthetic samples of the "
         "very rare 'disgust' and 'fear' classes blur the class boundaries. All networks are already trained with "
         "a class-weighted loss, which addresses the imbalance without synthetic data.")])
    doc.h("4.2 Overfitting / underfitting check", 2)
    doc.p("For every protocol run the gap between training and validation accuracy was recorded (the 'Fit check' "
          "column). Early stopping on the validation split and the L1 + L2 penalties keep the large datasets in "
          "the 'good fit' range. Where 'overfitting' is reported (mainly WESAD, with only 15 subjects), training "
          "accuracy is much higher than validation on an unseen person; early stopping keeps the weights of the "
          "best validation epoch, so the reported test scores are not affected by the later over-fitted epochs.")
    doc.h("4.3 Limitations", 2)
    doc.bullets([
        "Speech: clip-level split with overlapping speakers; a speaker-independent test (e.g. with CREMA-D's 91 "
        "actors) would give lower, more realistic scores.",
        "DeBERTa on GoEmotions was evaluated after 2 of 4 epochs because the run was interrupted.",
        "The 4 GB GPU limited batch sizes, so the heaviest models (ResNet50+ViT, DeBERTa) used small batches.",
        "WESAD has only 15 subjects; per-subject variability is large.",
        "Protocol results on deep features measure the separability of each model's learned representation; the "
        "held-out test results measure the complete trained model. Both are reported for transparency."])

    doc.h("5. Conclusion", 1)
    doc.p("All twenty hybrid architectures specified for the project were implemented, trained and compared on "
          "the same data splits with the same metrics, and evaluated with the prescribed protocol (80:20 / 70:30 "
          "hold-out and 10-fold cross-validation, each with and without SMOTE, with early stopping and L1/L2 "
          "regularisation). The best model of each modality (Section 3) is selected for the project's web "
          "application for real-time multimodal emotion recognition.")

    path = OUT / "Final_Results_Multimodal_Emotion.docx"
    doc.save(path)
    print(f"saved {path} ({doc.fig_no} figures, {doc.tab_no} tables)")


if __name__ == "__main__":
    main()
